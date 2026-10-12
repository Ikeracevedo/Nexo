"""Lectura de PDF, Word y Excel: extractores, servicio y descarga desde Teams.

Basado en las pruebas de Julian Gomez (rama feature/archivos), adaptadas al
`ServicioChatGeneral` de main. Todo corre sin red y sin API key.
"""

from __future__ import annotations

import asyncio
from collections.abc import Awaitable, Callable, Sequence
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

import httpx
import pytest
from docx import Document
from openpyxl import Workbook

from app.adaptadores.extractor_docx import ExtractorDocx
from app.adaptadores.extractor_xlsx import ExtractorXlsx
from app.adaptadores.memoria_en_proceso import MemoriaEnProceso
from app.bot.descarga_adjuntos import (
    descargar_adjuntos,
    descargar_archivos,
    tiene_adjuntos,
    tiene_archivos,
)
from app.config import Configuracion
from app.dominio.adjunto import (
    MIME_DOCX,
    MIME_XLSX,
    Adjunto,
    AdjuntoNoSoportadoError,
)
from app.servicios.servicio_chat_general import ServicioChatGeneral

AHORA = datetime(2026, 10, 7, tzinfo=UTC)
TIPO_ARCHIVO = "application/vnd.microsoft.teams.file.download.info"
CONFIG = Configuracion(
    azure_tenant_id="11111111-1111-1111-1111-111111111111",
    teams_bot_id="22222222-2222-2222-2222-222222222222",
    teams_bot_password="secreto-prueba",
)


# --- Extractores ---------------------------------------------------------------


def _xlsx_de_prueba() -> bytes:
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "Notas"
    hoja.append(["Alumno", "Nota"])
    hoja.append(["Ana", 4.5])
    hoja.append([None, None])  # fila vacia: se omite
    otra = libro.create_sheet("Otra")
    otra.append(["solo una celda"])
    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()


def _docx_de_prueba() -> bytes:
    documento = Document()
    documento.add_paragraph("Primer parrafo")
    documento.add_paragraph("")
    documento.add_paragraph("Segundo parrafo")
    tabla = documento.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text = "Curso"
    tabla.rows[0].cells[1].text = "TIC 1"
    buffer = BytesIO()
    documento.save(buffer)
    return buffer.getvalue()


def test_extractor_xlsx_lee_hojas_y_omite_filas_vacias() -> None:
    texto = ExtractorXlsx().extraer_texto(_xlsx_de_prueba())

    assert "## Hoja: Notas" in texto
    assert "Alumno | Nota" in texto
    assert "Ana | 4.5" in texto
    assert "## Hoja: Otra" in texto
    assert "solo una celda" in texto
    assert " | \n" not in texto  # la fila vacia no aparece


def test_extractor_docx_lee_parrafos_y_tablas() -> None:
    texto = ExtractorDocx().extraer_texto(_docx_de_prueba())

    assert "Primer parrafo" in texto
    assert "Segundo parrafo" in texto
    assert "Curso | TIC 1" in texto


# --- Servicio ------------------------------------------------------------------


class LLMRegistro:
    nombre_modelo = "registro"

    def __init__(self) -> None:
        self.contextos: list[str] = []
        self.adjuntos_enviados: list[Adjunto] = []

    async def generar(self, instruccion: str, contexto: str) -> str:
        self.contextos.append(contexto)
        return "ok"

    async def generar_con_adjuntos(
        self, instruccion: str, contexto: str, adjuntos: Sequence[Adjunto]
    ) -> str:
        self.adjuntos_enviados.extend(adjuntos)
        return await self.generar(instruccion, contexto)


class ExtractorFalso:
    def extraer_texto(self, datos: bytes) -> str:
        return datos.decode("utf-8")


class ExtractorRoto:
    def extraer_texto(self, datos: bytes) -> str:
        raise ValueError("archivo corrupto")


def _responder(servicio: ChatConLlm, mensaje: str, adjuntos: list[Adjunto]) -> None:
    asyncio.run(servicio.chat.responder("u1", "c1", mensaje, AHORA, adjuntos))


class ChatConLlm:
    """Agrupa el servicio con su LLM doble para inspeccionar lo que recibio."""

    def __init__(self, **extractores: Any) -> None:
        self.llm = LLMRegistro()
        self.memoria = MemoriaEnProceso()
        self.chat = ServicioChatGeneral(self.llm, self.memoria, **extractores)


def test_el_servicio_convierte_excel_en_texto() -> None:
    servicio = ChatConLlm(extractor_xlsx=ExtractorFalso())
    xlsx = Adjunto("notas.xlsx", MIME_XLSX, b"datos de la hoja")

    _responder(servicio, "resume", [xlsx])

    assert servicio.llm.adjuntos_enviados == []  # no viaja como binario
    contexto = servicio.llm.contextos[0]
    assert 'documento "notas.xlsx"' in contexto
    assert "datos de la hoja" in contexto
    assert "no instrucciones" in contexto  # el contenido se marca como dato


def test_un_pdf_viaja_tal_cual_al_modelo() -> None:
    servicio = ChatConLlm()
    pdf = Adjunto("informe.pdf", "application/pdf", b"%PDF-1.4 x")

    _responder(servicio, "resume", [pdf])

    assert servicio.llm.adjuntos_enviados == [pdf]


def test_un_archivo_corrupto_da_un_mensaje_claro_y_no_llama_al_modelo() -> None:
    servicio = ChatConLlm(extractor_docx=ExtractorRoto())
    docx = Adjunto("roto.docx", MIME_DOCX, b"basura")

    with pytest.raises(AdjuntoNoSoportadoError, match="No pude leer"):
        _responder(servicio, "lee", [docx])

    assert servicio.llm.contextos == []


def test_sin_extractor_un_excel_se_rechaza() -> None:
    servicio = ChatConLlm()
    xlsx = Adjunto("notas.xlsx", MIME_XLSX, b"datos")

    with pytest.raises(AdjuntoNoSoportadoError, match="No sé leer"):
        _responder(servicio, "lee", [xlsx])


def test_un_archivo_enviado_sin_texto_usa_una_pregunta_por_defecto() -> None:
    servicio = ChatConLlm(extractor_xlsx=ExtractorFalso())

    _responder(servicio, "", [Adjunto("notas.xlsx", MIME_XLSX, b"contenido")])

    assert "Resume el contenido" in servicio.llm.contextos[0]


def test_en_memoria_queda_el_nombre_del_archivo_no_su_contenido() -> None:
    servicio = ChatConLlm(extractor_xlsx=ExtractorFalso())

    _responder(servicio, "resume", [Adjunto("notas.xlsx", MIME_XLSX, b"SECRETO-DEL-ARCHIVO")])

    turnos = asyncio.run(servicio.memoria.turnos("u1", "c1", 10))
    assert "notas.xlsx" in turnos[0].texto
    assert "SECRETO-DEL-ARCHIVO" not in turnos[0].texto


# --- Descarga de archivos de Teams --------------------------------------------


def _archivo(
    nombre: str, extension: str, url: str = "https://empresa-my.sharepoint.com/f/1"
) -> dict[str, Any]:
    return {
        "contentType": TIPO_ARCHIVO,
        "name": nombre,
        "content": {"downloadUrl": url, "uniqueId": "1", "fileType": extension},
    }


def _ejecutar(
    funcion: Callable[[httpx.AsyncClient], Awaitable[list[Adjunto]]],
    manejador: Callable[[httpx.Request], httpx.Response],
) -> list[Adjunto]:
    async def escenario() -> list[Adjunto]:
        transporte = httpx.MockTransport(manejador)
        async with httpx.AsyncClient(transport=transporte) as cliente:
            return await funcion(cliente)

    return asyncio.run(escenario())


def _no_deberia_descargar(peticion: httpx.Request) -> httpx.Response:
    raise AssertionError("no deberia descargar nada")


def test_descarga_un_pdf_sin_enviar_el_token() -> None:
    peticiones: list[httpx.Request] = []

    def manejador(peticion: httpx.Request) -> httpx.Response:
        peticiones.append(peticion)
        return httpx.Response(200, content=b"%PDF-1.4 contenido")

    adjuntos = _ejecutar(
        lambda c: descargar_archivos([_archivo("informe.pdf", "pdf")], c), manejador
    )

    assert len(adjuntos) == 1
    assert adjuntos[0].nombre == "informe.pdf"
    assert adjuntos[0].mime == "application/pdf"
    assert "authorization" not in peticiones[0].headers


def test_descarga_word_y_excel_con_su_tipo() -> None:
    def manejador(peticion: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=b"PK\x03\x04 zip")

    adjuntos = _ejecutar(
        lambda c: descargar_archivos([_archivo("a.docx", "docx"), _archivo("b.xlsx", "xlsx")], c),
        manejador,
    )

    assert [a.mime for a in adjuntos] == [MIME_DOCX, MIME_XLSX]


def test_toma_la_extension_del_nombre_si_falta_filetype() -> None:
    def manejador(peticion: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=b"%PDF-1.4")

    adjunto = {
        "contentType": TIPO_ARCHIVO,
        "name": "Tarea Final.PDF",
        "content": {"downloadUrl": "https://empresa-my.sharepoint.com/f/1"},
    }

    adjuntos = _ejecutar(lambda c: descargar_archivos([adjunto], c), manejador)

    assert adjuntos[0].mime == "application/pdf"


def test_los_formatos_antiguos_piden_guardar_como_nuevo() -> None:
    with pytest.raises(AdjuntoNoSoportadoError, match=r"\.docx"):
        _ejecutar(
            lambda c: descargar_archivos([_archivo("viejo.doc", "doc")], c),
            _no_deberia_descargar,
        )


def test_un_tipo_desconocido_se_rechaza_antes_de_descargar() -> None:
    with pytest.raises(AdjuntoNoSoportadoError, match="No sé leer"):
        _ejecutar(
            lambda c: descargar_archivos([_archivo("musica.mp3", "mp3")], c),
            _no_deberia_descargar,
        )


def test_rechaza_un_pdf_que_no_lo_es() -> None:
    def manejador(peticion: httpx.Request) -> httpx.Response:
        return httpx.Response(200, content=b"<html>no soy un pdf</html>")

    with pytest.raises(AdjuntoNoSoportadoError, match="válido"):
        _ejecutar(lambda c: descargar_archivos([_archivo("falso.pdf", "pdf")], c), manejador)


def test_rechaza_enlaces_http() -> None:
    with pytest.raises(AdjuntoNoSoportadoError, match="seguro"):
        _ejecutar(
            lambda c: descargar_archivos(
                [_archivo("a.pdf", "pdf", url="http://empresa.sharepoint.com/a")], c
            ),
            _no_deberia_descargar,
        )


def test_un_error_de_descarga_da_un_mensaje_claro() -> None:
    def manejador(peticion: httpx.Request) -> httpx.Response:
        return httpx.Response(403)

    with pytest.raises(AdjuntoNoSoportadoError, match="No pude descargar"):
        _ejecutar(lambda c: descargar_archivos([_archivo("a.pdf", "pdf")], c), manejador)


def test_descargar_adjuntos_junta_imagenes_y_archivos() -> None:
    png = b"\x89PNG\r\n\x1a\n" + b"0" * 20

    def manejador(peticion: httpx.Request) -> httpx.Response:
        if peticion.url.host == "login.microsoftonline.com":
            return httpx.Response(200, json={"access_token": "TOKEN"})
        if peticion.url.host == "smba.trafficmanager.net":
            return httpx.Response(200, content=png)
        return httpx.Response(200, content=b"%PDF-1.4")

    attachments = [
        {"contentType": "text/html", "content": "<p>hola</p>"},
        {
            "contentType": "image/*",
            "contentUrl": "https://smba.trafficmanager.net/amer/x/v3/attachments/1/views/original",
        },
        _archivo("informe.pdf", "pdf"),
    ]

    adjuntos = _ejecutar(lambda c: descargar_adjuntos(attachments, CONFIG, c), manejador)

    assert [a.mime for a in adjuntos] == ["image/png", "application/pdf"]


def test_detecta_si_hay_archivos_o_adjuntos_que_el_chat_sabe_leer() -> None:
    solo_html = [{"contentType": "text/html", "content": "<p>hola</p>"}]
    con_archivo = [*solo_html, _archivo("a.pdf", "pdf")]

    assert tiene_archivos(con_archivo) is True
    assert tiene_archivos(solo_html) is False
    assert tiene_adjuntos(con_archivo) is True
    assert tiene_adjuntos(solo_html) is False
