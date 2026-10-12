from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

MAX_HOJAS = 5
MAX_FILAS_POR_HOJA = 200
MAX_COLUMNAS = 30


class ExtractorXlsx:
    """Convierte un .xlsx en texto: una seccion por hoja, celdas separadas por ' | '.

    Limites: primeras 5 hojas, 200 filas y 30 columnas por hoja, para no
    mandar al modelo hojas enormes.

    Limitaciones: lee los valores que Excel dejo calculados. Un archivo
    generado por un programa que nunca lo abrio en Excel puede traer las
    formulas sin resultado. No soporta el formato .xls antiguo y ignora
    graficos e imagenes.
    """

    def extraer_texto(self, datos: bytes) -> str:
        libro = load_workbook(BytesIO(datos), read_only=True, data_only=True)
        try:
            partes: list[str] = []

            for hoja in libro.worksheets[:MAX_HOJAS]:
                partes.append(f"## Hoja: {hoja.title}")

                for numero, fila in enumerate(hoja.iter_rows(values_only=True)):
                    if numero >= MAX_FILAS_POR_HOJA:
                        partes.append("[...filas omitidas]")
                        break

                    celdas = ["" if c is None else str(c) for c in fila[:MAX_COLUMNAS]]
                    if any(c.strip() for c in celdas):
                        partes.append(" | ".join(celdas))

            if len(libro.worksheets) > MAX_HOJAS:
                partes.append("[...hojas omitidas]")

            return "\n".join(partes)
        finally:
            libro.close()
