from __future__ import annotations

from io import BytesIO

from docx import Document


class ExtractorDocx:
    """Extrae el texto de un .docx (parrafos y tablas).

    Limitaciones: no soporta el formato .doc antiguo, ignora imagenes
    incrustadas y las tablas se anexan despues de los parrafos.
    """

    def extraer_texto(self, datos: bytes) -> str:
        documento = Document(BytesIO(datos))

        partes = [p.text for p in documento.paragraphs if p.text.strip()]

        for tabla in documento.tables:
            for fila in tabla.rows:
                celdas = [c.text.strip() for c in fila.cells]
                if any(celdas):
                    partes.append(" | ".join(celdas))

        return "\n".join(partes)
